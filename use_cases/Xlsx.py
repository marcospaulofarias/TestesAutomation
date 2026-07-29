from loguru import logger
import openpyxl # type: ignore
from use_cases.PrintAutomation import PrintAutomation
from os import path

class Xlsx:
    def __init__(self, file_path: str):
        """Classe para manipulação de arquivos Excel (.xlsx) usando openpyxl.

        :param file_path: Caminho completo do arquivo Excel a ser manipulado.
        """
        self.file_path = file_path
        if not path.exists(file_path):
            self.workbook = openpyxl.Workbook()
            self._save_workbook()
        self.workbook = openpyxl.load_workbook(file_path)
        self.printautomation = PrintAutomation()

    def get_sheet_names(self):
        """Retorna uma lista com os nomes das planilhas no arquivo Excel.
        
        :return: Lista com os nomes das planilhas no arquivo.
        """
        return self.workbook.sheetnames
    
    def get_sheet(self, sheet_name: str, create: bool=False):
        """Retorna a planilha especificada pelo nome.

        :param sheet_name: Nome da planilha a ser retornada.
        :param create: Criar planilha se não existir?
        :return: Objeto da planilha.
        """
        if sheet_name not in self.workbook.sheetnames:
            if create:
                self.workbook.create_sheet(sheet_name)
                return self.workbook[sheet_name]
            raise ValueError(f"Sheet '{sheet_name}' does not exist in the workbook.")
        return self.workbook[sheet_name]
    
    def get_cell_value(self, sheet_name: str, cell_reference: str):
        """Retorna o valor da célula especificada.

        :param sheet_name: Nome da planilha que contém a célula.
        :param cell_reference: Referência da célula (ex: 'A1').
        :return: Valor da célula.
        """
        sheet = self.get_sheet(sheet_name)
        return sheet[cell_reference].value
    
    def set_cell_value(self, sheet_name: str, cell_reference: str, value: any, create: bool=False):
        """Define o valor da célula especificada.

        :param sheet_name: Nome da planilha que contém a célula.
        :param cell_reference: Referência da célula (ex: 'A1').
        :param value: Valor a ser definido na célula.
        :param create: Criar planilha se não existir?
        """
        sheet = self.get_sheet(sheet_name, create=create)
        sheet[cell_reference].value = value
        self._save_workbook()

    def _save_workbook(self):
        try:
            self.workbook.save(self.file_path)
        except Exception as error_x:
            logger.critical(f'Erro ao salvar o arquivo "{self.file_path}": {error_x}')
            self.printautomation.print_error()
            raise RuntimeError(f'Erro ao salvar o arquivo "{self.file_path}": {error_x}') from error_x
